import csv
import json
import re
import uuid
from datetime import datetime, timedelta, timezone
from io import BytesIO, StringIO
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.platform.audit.model import AuditLog
from app.modules.organization.service import read_for_reporting as read_organizations
from app.modules.people_access.service import read_for_reporting as read_people
from app.core.datetime import parse_instant
from app.modules.record.service import read_for_reporting as read_records
from app.modules.storage_integration.service import read_for_reporting as read_storage, storage

RESOURCE_MAP = {
    "meetingTopics": "meeting-topics",
    "meetingHistory": "meeting-history",
    "incomingDocuments": "incoming-documents",
    "outgoingDocuments": "outgoing-documents",
    "documents": "documents",
    "masterListRequests": "master-list-requests",
    "recordLogs": "record-logs",
    "departments": "departments",
    "companies": "companies",
    "purposes": "purposes",
    "sectors": "sectors",
    "officers": "officers",
    "roles": "roles",
    "users": "users",
    "recordTypes": "record-types",
    "recordAttributes": "record-attributes",
    "fileUploads": "file-uploads",
    "googleDriveSync": "google-drive-sync",
    "portalLogs": "portal-logs",
    "systemLogs": "system-logs",
}


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return str(value)


def _parse_ids(selected: list) -> list[uuid.UUID]:
    ids: list[uuid.UUID] = []
    for raw in selected[:5000]:
        try:
            ids.append(uuid.UUID(str(raw)))
        except ValueError:
            pass
    return ids


SECRET_FIELDS = {"password", "passwordHash", "secret", "token", "objectKey"}


def date_bounds(payload: dict) -> tuple[datetime | None, datetime | None]:
    """Dialog startDate/endDate. Records filter on record_time; other resources on created_at."""
    start = parse_instant(str(payload.get("startDate") or "") or None)
    end = parse_instant(str(payload.get("endDate") or "") or None, end_of_day=True)
    return start, end


def range_filters(column, start, end):
    filters = []
    if start:
        filters.append(column >= start)
    if end:
        filters.append(column <= end)
    return filters


def select_fields(field_codes: list, values_list: list[dict]) -> list[str]:
    fields = [str(code) for code in field_codes or [] if str(code) not in SECRET_FIELDS]
    if not fields:
        fields = sorted({key for row in values_list for key in row.keys() if key not in SECRET_FIELDS})
    return ["id", *[f for f in fields if f != "id"]]


async def _rows_for_resource(db: AsyncSession, resource: str, payload: dict) -> list[dict]:
    selected = payload.get("selectedIds") or []
    ids = _parse_ids(selected) if payload.get("scope") == "selected" else None
    start, end = date_bounds(payload)

    for reader in (read_records, read_organizations, read_people, read_storage):
        rows = await reader(db, resource, ids, start, end)
        if rows is not None:
            return rows

    if resource in {"record-logs", "portal-logs", "system-logs"}:
        source = {"portal-logs": "portal", "system-logs": "system"}.get(resource, "record")
        filters = (
            [AuditLog.source_log == source]
            if source != "record"
            else [AuditLog.source_log.in_(("record", "api", "unknown"))]
        )
        filters.extend(range_filters(AuditLog.created_at, start, end))
        rows = (
            await db.scalars(
                select(AuditLog)
                .where(*filters)
                .order_by(AuditLog.created_at.desc())
                .limit(10000)
            )
        ).all()
        return [
            {
                "id": str(row.id),
                "summary": row.message or row.action_code,
                "action": row.action_code,
                "entityType": row.table_name,
                "occurredAt": row.created_at.isoformat() if row.created_at else None,
                "status": row.status_code,
            }
            for row in rows
        ]
    return []

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXPORT_TEMPLATE_DIR = Path(__file__).resolve().parents[3] / "resources" / "export_templates"


def _header_key(value) -> str:
    """Normalize a header/label for matching (case/space/punctuation-insensitive)."""
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def _template_path(resource: str) -> Path:
    return EXPORT_TEMPLATE_DIR / f"{resource}.xlsx"


def _xlsx_bytes(resource: str, fields: list[str], labels: dict, values_list: list[dict]) -> bytes:
    """Fill the resource Excel template, or build a plain sheet when none exists.

    Template contract: first sheet, header row 1. Each header cell is matched to a
    field by its localized label (or field code); matched columns get data rows
    appended from row 2. Unmatched columns are left blank.
    """
    from openpyxl import Workbook, load_workbook

    template = _template_path(resource)
    if template.exists():
        workbook = load_workbook(template)
        sheet = workbook.active
        lookup: dict[str, str] = {}
        for code in fields:
            lookup[_header_key(code)] = code
            label = labels.get(code)
            if label:
                lookup[_header_key(label)] = code
        header = next(sheet.iter_rows(min_row=1, max_row=1))
        column_field = {
            index: lookup[key]
            for index, cell in enumerate(header, start=1)
            if (key := _header_key(cell.value)) in lookup
        }
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = (resource[:31] or "export")
        column_field = {}
        for index, code in enumerate(fields, start=1):
            sheet.cell(row=1, column=index, value=labels.get(code) or code)
            column_field[index] = code
    for offset, values in enumerate(values_list):
        row_index = 2 + offset
        for column, code in column_field.items():
            sheet.cell(row=row_index, column=column, value=_cell(values.get(code)))
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def generate_export(db: AsyncSession, job: object) -> None:
    payload = dict(job.payload or {})
    resource = RESOURCE_MAP.get(str(payload.get("resource")), str(payload.get("resource") or ""))
    if not resource:
        raise ValueError("Export resource is required")
    values_list = await _rows_for_resource(db, resource, payload)
    fields = select_fields(list(payload.get("fieldCodes") or []), values_list)
    labels = {str(key): str(value) for key, value in (payload.get("fieldLabels") or {}).items()}
    stamp = datetime.now(timezone.utc)
    if str(payload.get("format") or "csv").lower() == "xlsx":
        data = _xlsx_bytes(resource, fields, labels, values_list)
        filename = f"{resource}-{stamp:%Y%m%d-%H%M%S}.xlsx"
        mime_type = XLSX_MIME
    else:
        output = StringIO(newline="")
        writer = csv.DictWriter(output, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for values in values_list:
            writer.writerow({key: _cell(values.get(key)) for key in fields})
        data = output.getvalue().encode("utf-8-sig")
        filename = f"{resource}-{stamp:%Y%m%d-%H%M%S}.csv"
        mime_type = "text/csv; charset=utf-8"
    key = f"exports/{job.id}/{filename}"
    await storage.put_bytes(key, data, mime_type)
    payload.update({
        "status": "completed",
        "objectKey": key,
        "fileName": filename,
        "mimeType": mime_type,
        "sizeBytes": len(data),
        "downloadUrl": f"/api/v2/files/{job.id}",
        "completedAt": stamp.isoformat(),
        "expiresAt": (stamp + timedelta(hours=24)).isoformat(),
    })
    job.payload = payload
